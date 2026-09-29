"""
Генератор синтетических тестовых документов, expected.json и docs/RULES.md.

Все имена, реквизиты и организации вымышлены. ИНН начинается с кода региона 00,
которого не существует, поэтому совпасть с настоящей организацией не может.

Запуск из корня репозитория:
    pip install -r testdata/requirements.txt
    python testdata/generate.py           # перегенерировать файлы
    python testdata/generate.py --check   # проверить, что файлы в репозитории актуальны (CI)

Генерация воспроизводима побайтно: шрифт лежит в testdata/fonts, все зависимости,
включая транзитивные, зафиксированы в requirements.txt, время внутри DOCX
фиксировано, а DOCX и PDF пишутся без сжатия — байты не зависят от версии zlib.
--check сравнивает файлы побайтно.

Источник истины для замечаний — Java-тест TestdataRulesTest с настоящим
DefaultRuleEngine. Здесь лишь предварительная проверка подмножества его логики
(REQUIRED, dd.MM.uuuu, совпадение всей строки с шаблоном, длина, ONE_OF),
чтобы ошибка в данных всплывала сразу при генерации.
"""
import json
import re
import sys
import tempfile
import zipfile
from datetime import date, datetime
from pathlib import Path

from docx import Document
from docx.shared import Pt
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

ROOT = Path(__file__).resolve().parent.parent
TESTDATA = ROOT / "testdata"
FONT = TESTDATA / "fonts" / "DejaVuSans.ttf"
RULES = json.loads((TESTDATA / "rules.json").read_text(encoding="utf-8"))

# «Сегодня» для DATE_NOT_FUTURE. Тесты движка передают её в ValidationContext.
REFERENCE_DATE = date(2026, 9, 21)
FIXED_TIME = datetime(2026, 9, 21, 12, 0, 0)
FOOTER = "Синтетический тестовый документ. Все имена, организации и реквизиты вымышлены."


class GenerationError(Exception):
    """Данные противоречат правилам или файлы в репозитории устарели."""


# ---------------------------------------------------------------------------
# Документы. lines — строки по порядку; fields — результат идеального извлечения;
# normalized — поля, чьё значение приведено к нужной форме и дословно в тексте
# отсутствует (падеж); issues — коды правил, которые обязаны сработать.
# ---------------------------------------------------------------------------

MEMO_BODY = (
    "В связи с расширением отдела продаж с 1 октября 2026 года прошу согласовать "
    "закупку пяти ноутбуков для новых сотрудников. Ориентировочная стоимость — "
    "450 000 рублей, средства предусмотрены бюджетом отдела на IV квартал."
)
VACATION_HEAD = [
    ("right", "Генеральному директору ООО «Демо-компания» Орлову В.П."),
    ("right", "от менеджера по продажам Кузнецова Игоря Олеговича"),
    ("", ""),
    ("bold", "ЗАЯВЛЕНИЕ"),
    ("", ""),
]
VACATION_NORMALIZED = ["employee_name", "employee_position"]
TRIP_HEAD = [
    ("right", "Генеральному директору ООО «Демо-компания» Орлову В.П."),
    ("right", "от менеджера по продажам Кузнецова Игоря Олеговича"),
    ("", ""),
    ("bold", "ЗАЯВКА НА КОМАНДИРОВКУ"),
    ("", ""),
]
TRIP_NORMALIZED = ["employee_name"]

DOCS = [
    {
        "file": "01_memo_ok.docx",
        "type": "OFFICIAL_MEMO",
        "about": "Служебная записка без замечаний.",
        "lines": [
            ("right", "Генеральному директору ООО «Демо-компания» Орлову В.П."),
            ("", ""),
            ("bold", "СЛУЖЕБНАЯ ЗАПИСКА"),
            ("", "15.09.2026 № СЗ-117"),
            ("", ""),
            ("", "О закупке ноутбуков для отдела продаж"),
            ("", ""),
            ("", MEMO_BODY),
            ("", ""),
            ("", "Руководитель отдела продаж"),
            ("", "А.С. Петрова"),
        ],
        "fields": {
            "addressee": "Генеральному директору ООО «Демо-компания» Орлову В.П.",
            "author_name": "А.С. Петрова",
            "doc_date": "15.09.2026",
            "reg_number": "СЗ-117",
            "subject": "О закупке ноутбуков для отдела продаж",
            "body": MEMO_BODY,
            "signer_position": "Руководитель отдела продаж",
            "signer_name": "А.С. Петрова",
        },
        "normalized": [],
        "issues": [],
    },
    {
        "file": "02_memo_errors.docx",
        "type": "OFFICIAL_MEMO",
        "about": "Нет адресата, дата в будущем, заголовок без «О», в подписи полное имя.",
        "lines": [
            ("bold", "СЛУЖЕБНАЯ ЗАПИСКА"),
            ("", "31.12.2099 № СЗ-118"),
            ("", ""),
            ("", "Закупка ноутбуков для отдела продаж"),
            ("", ""),
            ("", MEMO_BODY),
            ("", ""),
            ("", "Руководитель отдела продаж"),
            ("", "Анна Петрова"),
        ],
        "fields": {
            "author_name": "Анна Петрова",
            "doc_date": "31.12.2099",
            "reg_number": "СЗ-118",
            "subject": "Закупка ноутбуков для отдела продаж",
            "body": MEMO_BODY,
            "signer_position": "Руководитель отдела продаж",
            "signer_name": "Анна Петрова",
        },
        "normalized": [],
        "issues": ["M1", "M4", "M6", "M10"],
    },
    {
        "file": "03_memo_word_date.pdf",
        "type": "OFFICIAL_MEMO",
        "about": (
            "PDF. Дата без года — не понять, о каком годе речь (словесно-цифровой способ с годом ГОСТ допускает). "
            "Номер без индекса, короткий текст, в подписи нет должности."
        ),
        "lines": [
            ("right", "Главному бухгалтеру ООО «Демо-компания» Соколовой Е.Н."),
            ("", ""),
            ("bold", "СЛУЖЕБНАЯ ЗАПИСКА"),
            ("", "15 сентября № 119"),
            ("", ""),
            ("", "О возмещении командировочных расходов"),
            ("", ""),
            ("", "Прошу возместить расходы."),
            ("", ""),
            ("", "Д.В. Смирнов"),
        ],
        "fields": {
            "addressee": "Главному бухгалтеру ООО «Демо-компания» Соколовой Е.Н.",
            "author_name": "Д.В. Смирнов",
            "doc_date": "15 сентября",
            "reg_number": "119",
            "subject": "О возмещении командировочных расходов",
            "body": "Прошу возместить расходы.",
            "signer_name": "Д.В. Смирнов",
        },
        "normalized": [],
        "issues": ["M3", "M5", "M8", "M9"],
    },
    {
        "file": "04_vacation_ok.docx",
        "type": "VACATION_REQUEST",
        "about": "Заявление на отпуск без замечаний.",
        "lines": VACATION_HEAD + [
            ("", "Прошу предоставить мне ежегодный оплачиваемый отпуск с 01.10.2026 "
                 "продолжительностью 14 календарных дней."),
            ("", ""),
            ("", "15.09.2026        И.О. Кузнецов"),
        ],
        "fields": {
            "employee_name": "Кузнецов Игорь Олегович",
            "employee_position": "менеджер по продажам",
            "vacation_type": "ежегодный оплачиваемый",
            "start_date": "01.10.2026",
            "days_count": "14",
        },
        "normalized": VACATION_NORMALIZED,
        "issues": [],
    },
    {
        "file": "05_vacation_errors.pdf",
        "type": "VACATION_REQUEST",
        "about": "PDF. Учебный отпуск, дата начала без года, дни словами.",
        "lines": VACATION_HEAD + [
            ("", "Прошу предоставить мне учебный отпуск с 1 октября "
                 "продолжительностью двадцать восемь календарных дней."),
            ("", ""),
            ("", "15.09.2026        И.О. Кузнецов"),
        ],
        "fields": {
            "employee_name": "Кузнецов Игорь Олегович",
            "employee_position": "менеджер по продажам",
            "vacation_type": "учебный",
            "start_date": "1 октября",
            "days_count": "двадцать восемь",
        },
        "normalized": VACATION_NORMALIZED,
        "issues": ["V2", "V4", "V5"],
    },
    {
        "file": "06_support_ok.pdf",
        "type": "SUPPORT_MEASURE_REQUEST",
        "about": "PDF. Заявка на меру поддержки без замечаний.",
        "lines": [
            ("right", "В центр поддержки предпринимательства (вымышленный)"),
            ("", ""),
            ("bold", "ЗАЯВКА НА ПОЛУЧЕНИЕ МЕРЫ ПОДДЕРЖКИ"),
            ("", ""),
            ("", "Заявитель: ООО «Демо-компания»"),
            ("", "ИНН: 0012345678"),
            ("", "Категория субъекта МСП: малое предприятие"),
            ("", "Вид поддержки: консультационная поддержка по вопросам бухгалтерского учёта"),
            ("", ""),
            ("", "Генеральный директор        В.П. Орлов        15.09.2026"),
        ],
        "fields": {
            "applicant_name": "ООО «Демо-компания»",
            "applicant_inn": "0012345678",
            "smb_category": "малое предприятие",
            "support_type": "консультационная поддержка по вопросам бухгалтерского учёта",
        },
        "normalized": [],
        "issues": [],
    },
    {
        "file": "07_support_errors.docx",
        "type": "SUPPORT_MEASURE_REQUEST",
        "about": "Не указан вид поддержки, ИНН из 9 цифр, форма вместо категории МСП.",
        "lines": [
            ("right", "В центр поддержки предпринимательства (вымышленный)"),
            ("", ""),
            ("bold", "ЗАЯВКА НА ПОЛУЧЕНИЕ МЕРЫ ПОДДЕРЖКИ"),
            ("", ""),
            ("", "Заявитель: ООО «Демо-компания»"),
            ("", "ИНН: 001234567"),
            ("", "Категория субъекта МСП: индивидуальный предприниматель"),
            ("", ""),
            ("", "Генеральный директор        В.П. Орлов        15.09.2026"),
        ],
        "fields": {
            "applicant_name": "ООО «Демо-компания»",
            "applicant_inn": "001234567",
            "smb_category": "индивидуальный предприниматель",
        },
        "normalized": [],
        "issues": ["S2", "S4", "S6"],
    },
    {
        "file": "08_memo_missing.pdf",
        "type": "OFFICIAL_MEMO",
        "about": "PDF. Нет даты и нет текста записки: пустые поля ловят только правила «Заполнено».",
        "lines": [
            ("right", "Генеральному директору ООО «Демо-компания» Орлову В.П."),
            ("", ""),
            ("bold", "СЛУЖЕБНАЯ ЗАПИСКА"),
            ("", "№ СЗ-120"),
            ("", ""),
            ("", "О закупке ноутбуков для отдела продаж"),
            ("", ""),
            ("", "Руководитель отдела продаж"),
            ("", "А.С. Петрова"),
        ],
        "fields": {
            "addressee": "Генеральному директору ООО «Демо-компания» Орлову В.П.",
            "author_name": "А.С. Петрова",
            "reg_number": "СЗ-120",
            "subject": "О закупке ноутбуков для отдела продаж",
            "signer_position": "Руководитель отдела продаж",
            "signer_name": "А.С. Петрова",
        },
        "normalized": [],
        "issues": ["M2", "M7"],
    },
    {
        "file": "09_vacation_missing.docx",
        "type": "VACATION_REQUEST",
        "about": "Не указаны сотрудник и дата начала отпуска.",
        "lines": [
            ("right", "Генеральному директору ООО «Демо-компания» Орлову В.П."),
            ("", ""),
            ("bold", "ЗАЯВЛЕНИЕ"),
            ("", ""),
            ("", "Прошу предоставить мне ежегодный оплачиваемый отпуск "
                 "продолжительностью 14 календарных дней."),
            ("", ""),
            ("", "15.09.2026"),
        ],
        "fields": {
            "vacation_type": "ежегодный оплачиваемый",
            "days_count": "14",
        },
        "normalized": [],
        "issues": ["V1", "V3"],
    },
    {
        "file": "10_support_missing.pdf",
        "type": "SUPPORT_MEASURE_REQUEST",
        "about": "PDF. Не указаны заявитель, ИНН и категория МСП: пустые поля ловят только правила «Заполнено».",
        "lines": [
            ("right", "В центр поддержки предпринимательства (вымышленный)"),
            ("", ""),
            ("bold", "ЗАЯВКА НА ПОЛУЧЕНИЕ МЕРЫ ПОДДЕРЖКИ"),
            ("", ""),
            ("", "Вид поддержки: консультационная поддержка по вопросам бухгалтерского учёта"),
            ("", ""),
            ("", "Генеральный директор        В.П. Орлов        15.09.2026"),
        ],
        "fields": {
            "support_type": "консультационная поддержка по вопросам бухгалтерского учёта",
        },
        "normalized": [],
        "issues": ["S1", "S3", "S5"],
    },
    {
        "file": "11_order_other.docx",
        "type": "GENERIC",
        "about": (
            "Другой документ — приказ: общая проверка реквизитов. Дата словесно-цифровым способом, "
            "все реквизиты на месте — замечаний нет."
        ),
        "lines": [
            ("right", "ООО «Демо-компания»"),
            ("", ""),
            ("bold", "ПРИКАЗ"),
            ("", "«18» сентября 2026 г.        № 42-П"),
            ("", ""),
            ("", "О проведении инвентаризации"),
            ("", ""),
            ("", "В связи с окончанием квартала приказываю провести инвентаризацию основных средств "
                 "до 30 сентября 2026 года. Ответственным назначить главного бухгалтера Е.Н. Соколову."),
            ("", ""),
            ("", "Генеральный директор        В.П. Орлов"),
        ],
        "fields": {
            "doc_kind": "ПРИКАЗ",
            "doc_date": "«18» сентября 2026 г.",
            "reg_number": "42-П",
            "subject": "О проведении инвентаризации",
            "signer_position": "Генеральный директор",
            "signer_name": "В.П. Орлов",
        },
        "normalized": [],
        "issues": [],
    },
    {
        "file": "12_act_errors.pdf",
        "type": "GENERIC",
        "about": "PDF. Другой документ — акт без даты, номера и подписи с расшифровкой: только предупреждения.",
        "lines": [
            ("right", "ООО «Демо-компания»"),
            ("", ""),
            ("bold", "АКТ"),
            ("", "приёма-передачи оборудования"),
            ("", ""),
            ("", "Ноутбуки в количестве пяти штук переданы отделу продаж в исправном состоянии."),
            ("", ""),
            ("", "Передал: склад        Принял: отдел продаж"),
        ],
        "fields": {
            "doc_kind": "АКТ",
            "subject": "приёма-передачи оборудования",
        },
        "normalized": [],
        "issues": ["G1", "G4", "G5"],
    },
    {
        "file": "13_letter_errors.docx",
        "type": "GENERIC",
        "about": "Другой документ — письмо: дата без числа не читается, заголовка нет.",
        "lines": [
            ("right", "Директору ООО «Ромашка» (вымышленное)"),
            ("", ""),
            ("", "Исх. № 57 от сентябрь 2026 г."),
            ("", ""),
            ("", "Уважаемый коллега! Направляем график поставок на октябрь для согласования."),
            ("", ""),
            ("", "Руководитель отдела продаж        А.С. Петрова"),
        ],
        "fields": {
            "doc_date": "сентябрь 2026 г.",
            "reg_number": "57",
            "signer_position": "Руководитель отдела продаж",
            "signer_name": "А.С. Петрова",
        },
        "normalized": [],
        "issues": ["G2", "G3"],
    },
    {
        "file": "14_trip_ok.docx",
        "type": "BUSINESS_TRIP_REQUEST",
        "about": "Заявка на командировку без замечаний.",
        "lines": TRIP_HEAD + [
            ("", "Прошу направить меня в служебную командировку."),
            ("", ""),
            ("", "Место командировки: г. Казань"),
            ("", "Цель командировки: участие в выставке «Промэкспо»"),
            ("", "Дата начала: 05.10.2026"),
            ("", "Срок: 5 календарных дней"),
            ("", ""),
            ("", "15.09.2026        И.О. Кузнецов"),
        ],
        "fields": {
            "employee_name": "Кузнецов Игорь Олегович",
            "destination": "г. Казань",
            "purpose": "участие в выставке «Промэкспо»",
            "start_date": "05.10.2026",
            "days_count": "5",
        },
        "normalized": TRIP_NORMALIZED,
        "issues": [],
    },
    {
        "file": "15_trip_errors.pdf",
        "type": "BUSINESS_TRIP_REQUEST",
        "about": "PDF. Не указана цель, дата начала без года, срок словами.",
        "lines": TRIP_HEAD + [
            ("", "Прошу направить меня в служебную командировку."),
            ("", ""),
            ("", "Место командировки: г. Казань"),
            ("", "Дата начала: 5 октября"),
            ("", "Срок: пять календарных дней"),
            ("", ""),
            ("", "15.09.2026        И.О. Кузнецов"),
        ],
        "fields": {
            "employee_name": "Кузнецов Игорь Олегович",
            "destination": "г. Казань",
            "start_date": "5 октября",
            "days_count": "пять",
        },
        "normalized": TRIP_NORMALIZED,
        "issues": ["T3", "T5", "T6"],
    },
    {
        "file": "16_trip_missing.docx",
        "type": "BUSINESS_TRIP_REQUEST",
        "about": "Не указаны сотрудник, место и дата начала: пустые поля ловят только правила «Заполнено».",
        "lines": [
            ("right", "Генеральному директору ООО «Демо-компания» Орлову В.П."),
            ("", ""),
            ("bold", "ЗАЯВКА НА КОМАНДИРОВКУ"),
            ("", ""),
            ("", "Цель командировки: переговоры с заказчиком о графике поставок"),
            ("", "Срок: 3 календарных дня"),
            ("", ""),
            ("", "15.09.2026"),
        ],
        "fields": {
            "purpose": "переговоры с заказчиком о графике поставок",
            "days_count": "3",
        },
        "normalized": [],
        "issues": ["T1", "T2", "T4"],
    },
]


# ---------------------------------------------------------------------------
# Предварительная проверка данных
# ---------------------------------------------------------------------------

RU_MONTHS = {
    "января": 1, "февраля": 2, "марта": 3, "апреля": 4, "мая": 5, "июня": 6,
    "июля": 7, "августа": 8, "сентября": 9, "октября": 10, "ноября": 11, "декабря": 12,
}


def safe_date(year, month, day):
    try:
        return date(int(year), int(month), int(day))
    except ValueError:
        return None


def parse_ru_date(value: str):
    """Тот же алгоритм, что RussianDates.parse в бэкенде: цифровая, словесно-цифровая дата и ISO."""
    text = re.sub(r"[«»\"“”„‘’']", "", value).replace(" ", " ").lower()
    text = re.sub(r"\s+", " ", text).strip()
    text = re.sub(r"\s*(г\.?|год|года)\s*$", "", text).strip()
    if m := re.fullmatch(r"(\d{1,2})[./-](\d{1,2})[./-](\d{4})", text):
        return safe_date(m[3], m[2], m[1])
    if m := re.fullmatch(r"(\d{1,2})\s+([а-яё]+)\s+(\d{4})", text):
        month = RU_MONTHS.get(m[2])
        return safe_date(m[3], month, m[1]) if month else None
    if m := re.fullmatch(r"(\d{4})-(\d{2})-(\d{2})", text):
        return safe_date(m[1], m[2], m[3])
    return None


def parse_date(value: str, pattern: str):
    if pattern == "RU_DATE":
        return parse_ru_date(value)
    if pattern != "dd.MM.uuuu":
        raise GenerationError(f"Предпроверка знает только RU_DATE и dd.MM.uuuu, а не {pattern}")
    if not re.fullmatch(r"\d{2}\.\d{2}\.\d{4}", value):
        return None
    day, month, year = (int(part) for part in value.split("."))
    return safe_date(year, month, day)


def violates(rule: dict, value) -> bool:
    missing = value is None or value.strip() == ""
    check = rule["check"]
    if check == "REQUIRED":
        return missing
    if missing:
        return False
    if check == "DATE_FORMAT":
        return parse_date(value, rule["expected"]) is None
    if check == "DATE_NOT_FUTURE":
        # Неразборчивую дату это правило не судит: о формате сообщает DATE_FORMAT.
        parsed = parse_date(value, rule["expected"])
        return parsed is not None and parsed > REFERENCE_DATE
    if check == "MATCHES_PATTERN":
        return re.fullmatch(rule["expected"], value) is None
    if check == "MIN_LENGTH":
        return len(value) < int(rule["expected"])
    if check == "ONE_OF":
        return value not in json.loads(rule["expected"])
    raise GenerationError(f"Неизвестная проверка {check}")


def collapse(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def precheck() -> None:
    types = {t["code"]: t for t in RULES["types"]}
    for doc in DOCS:
        doc_type = types.get(doc["type"])
        if doc_type is None:
            raise GenerationError(f"{doc['file']}: тип {doc['type']} не описан в rules.json")
        known = {f["name"] for f in doc_type["fields"]}
        unknown = (set(doc["fields"]) | set(doc["normalized"])) - known
        if unknown:
            raise GenerationError(f"{doc['file']}: неизвестные поля {sorted(unknown)}")
        actual = [r["code"] for r in doc_type["rules"] if violates(r, doc["fields"].get(r["field"]))]
        if actual != doc["issues"]:
            raise GenerationError(f"{doc['file']}: ожидали {doc['issues']}, предпроверка даёт {actual}")
        text = collapse(" ".join(line for _, line in doc["lines"]))
        for name, value in doc["fields"].items():
            if name not in doc["normalized"] and collapse(value) not in text:
                raise GenerationError(f"{doc['file']}: значения поля {name} нет в тексте документа")


# ---------------------------------------------------------------------------
# Файлы
# ---------------------------------------------------------------------------

def write_docx(doc: dict, path: Path) -> None:
    document = Document()
    style = document.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    for kind, text in doc["lines"]:
        paragraph = document.add_paragraph()
        run = paragraph.add_run(text)
        if kind == "bold":
            run.bold = True
            paragraph.alignment = 1
        elif kind == "right":
            paragraph.alignment = 2
    document.add_paragraph()
    footer = document.add_paragraph().add_run(FOOTER)
    footer.italic = True
    footer.font.size = Pt(9)
    props = document.core_properties
    props.author = "Сибирский Вайб — тестовые данные"
    props.title = doc["about"]
    props.created = FIXED_TIME
    props.modified = FIXED_TIME
    props.last_modified_by = "generate.py"
    props.revision = 1
    document.save(path)
    fix_zip_times(path)


def fix_zip_times(path: Path) -> None:
    """Фиксированное время записей и хранение без сжатия: байты не зависят ни от часов, ни от версии zlib."""
    with zipfile.ZipFile(path) as source:
        entries = [(info.filename, source.read(info.filename)) for info in source.infolist()]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_STORED) as target:
        for name, data in entries:
            info = zipfile.ZipInfo(name, date_time=FIXED_TIME.timetuple()[:6])
            info.compress_type = zipfile.ZIP_STORED
            info.external_attr = 0o644 << 16
            # Иначе Python пишет в архив текущую ОС (0 — Windows, 3 — Unix), и байты на машинах расходятся.
            info.create_system = 3
            target.writestr(info, data)


def register_font() -> str:
    if not FONT.exists():
        raise GenerationError(f"Нет шрифта {FONT.relative_to(ROOT)}")
    pdfmetrics.registerFont(TTFont("DejaVuSans", str(FONT)))
    return "DejaVuSans"


def write_pdf(doc: dict, path: Path, font: str) -> None:
    base = ParagraphStyle("base", fontName=font, fontSize=11, leading=15)
    styles = {
        "": base,
        "bold": ParagraphStyle("bold", parent=base, alignment=1, fontSize=12),
        "right": ParagraphStyle("right", parent=base, alignment=2),
    }
    small = ParagraphStyle("small", parent=base, fontSize=8, leading=10)
    story = []
    for kind, text in doc["lines"]:
        story.append(Paragraph(text, styles[kind]) if text else Spacer(1, 6 * mm))
    story += [Spacer(1, 10 * mm), Paragraph(FOOTER, small)]
    pdf = SimpleDocTemplate(
        str(path), pagesize=A4, leftMargin=25 * mm, rightMargin=15 * mm,
        topMargin=20 * mm, bottomMargin=20 * mm, pageCompression=0,
        title=doc["about"], author="Сибирский Вайб — тестовые данные", invariant=1,
    )
    pdf.build(story)


def expected_json() -> str:
    expected = {
        "comment": (
            "Ожидаемый результат для каждого тестового документа. fields — результат идеального "
            "извлечения полей; normalized — поля, чьё значение приведено к нужной форме (падеж) "
            "и дословно в тексте отсутствует; issues — правила из rules.json, которые обязаны "
            "сработать, в порядке правил. referenceDate — «сегодня» для DATE_NOT_FUTURE."
        ),
        "referenceDate": REFERENCE_DATE.isoformat(),
        "documents": [
            {
                "file": f"documents/{d['file']}",
                "type": d["type"],
                "about": d["about"],
                "fields": d["fields"],
                "normalized": d["normalized"],
                "issues": d["issues"],
            }
            for d in DOCS
        ],
    }
    return json.dumps(expected, ensure_ascii=False, indent=2) + "\n"


# ---------------------------------------------------------------------------
# docs/RULES.md
# ---------------------------------------------------------------------------

KIND = {"LEGAL": "Закон", "INTERNAL_POLICY": "Правило компании", "PRODUCT_RULE": "Рекомендация сервиса"}
SEVERITY = {"BLOCKER": "Блокирует отправку", "WARNING": "Предупреждение", "INFO": "Подсказка"}
CHECK = {
    "REQUIRED": "Заполнено",
    "DATE_FORMAT": "Формат даты",
    "DATE_NOT_FUTURE": "Дата не в будущем",
    "MATCHES_PATTERN": "Соответствует шаблону",
    "MIN_LENGTH": "Не короче",
    "ONE_OF": "Одно из значений",
}
SHORT_SOURCE = {
    "GOST": "ГОСТ Р 7.0.97-2016",
    "MEMO_POLICY": "Типовые правила записок (шаблон)",
    "VACATION_POLICY": "Типовые правила отпусков (шаблон)",
    "TRIP_POLICY": "Типовые правила командировок (шаблон)",
    "PRODUCT": "Рекомендация сервиса",
    "LAW_209_ART_4_1": "209-ФЗ",
}


def cell(text) -> str:
    """Экранирует то, что ломает таблицу Markdown: «|» делит столбцы, «<» GitHub принимает за тег."""
    if text in (None, ""):
        return "—"
    return str(text).replace("|", "\\|").replace("<", "&lt;").replace("\n", " ")


def rules_md() -> str:
    sources = RULES["sources"]
    out = [
        "# Правила проверки документов",
        "",
        "Сгенерировано из [`testdata/rules.json`](../testdata/rules.json) скриптом",
        "`testdata/generate.py` — правьте JSON и перезапускайте скрипт, а не этот файл.",
        f"Источники сверены {datetime.fromisoformat(RULES['sourceCheckedAt']).strftime('%d.%m.%Y')}.",
        "",
        "Модель только извлекает поля, правила проверяет обычный код — `RuleEngine`",
        "(docs/DESIGN-DECISIONS.md, «Модель извлекает, код решает»). Пустое поле проверяет только",
        "«Заполнено»: остальные проверки его пропускают, иначе одно отсутствующее поле",
        "давало бы несколько одинаковых замечаний.",
        "",
        "«Дата не в будущем»: «сегодня» считается по часовому поясу компании — городу, который выбран при её",
        "создании и который администратор меняет в разделе «Компания». Дату ставит человек по своим часам: у компании",
        "в Новосибирске после полуночи уже новый день, пока в Москве ещё вчера.",
        "",
        "**Виды правил** — то, что пользователь и жюри видят у каждого замечания:",
        "- **Закон** — только когда закон сам определяет проверяемое требование. Если закон лишь",
        "  объясняет, зачем проверка нужна, это рекомендация сервиса или правило компании, а закон упомянут",
        "  в тексте замечания. Иначе продукт заявлял бы правовое основание, которого нет.",
        "- **Правило компании** — внутренний стандарт. Для коммерческой организации",
        "  ГОСТ Р 7.0.97-2016 добровольный, поэтому это правило компании, а не закон. Там, где",
        "  требование — выбор компании внутри рамок ГОСТ, источник — типовые правила сервиса.",
        "- **Рекомендация сервиса** — наше допущение, без которого проверка теряет смысл. Если компания",
        "  указала для неё свой документ-основание, она становится правилом компании.",
        "",
        "**Это шаблон, а не правила конкретной компании**. Компания получает его сразу, а",
        "администратор в разделе «Компания» → «Правила проверки» настраивает под свои локальные акты:",
        "выключить правило, сменить важность, текст и источник, а у «Не короче» и «Одно из значений» —",
        "ожидаемое значение. Правила вида «Закон» не выключаются и не меняются: закон одинаков для всех.",
        "Демо-компания при создании получает свои «Инструкцию по делопроизводству» и «Положение об",
        "отпусках» — так настройка видна в демонстрации сразу.",
        "",
    ]
    for doc_type in RULES["types"]:
        out += [f"## {doc_type['name']} (`{doc_type['code']}`)", "", "### Поля", "",
                "| Поле | Тип | Название | Подсказка для извлечения |", "|---|---|---|---|"]
        for f in doc_type["fields"]:
            out.append(f"| `{f['name']}` | {f['type']} | {cell(f['label'])} | {cell(f['hint'])} |")
        out += ["", "### Правила", "",
                "| Код | Поле | Проверка | Параметр | Вид | Важность | Источник | Текст замечания |",
                "|---|---|---|---|---|---|---|---|"]
        for r in doc_type["rules"]:
            src = sources[r["source"]]
            ref = f"{SHORT_SOURCE[r['source']]}, {r['sourceRef']}" if r["sourceRef"] else SHORT_SOURCE[r["source"]]
            link = f"[{cell(ref)}]({src['url']})" if src["url"] else cell(ref)
            # Внутри кода в таблице GitHub понимает только экранированный «\|».
            if r["expected"] == "RU_DATE":
                param = "полная дата: 23.09.2026 или «23» сентября 2026 г."
            else:
                param = "`" + r["expected"].replace("|", "\\|") + "`" if r["expected"] else "—"
            out.append(
                f"| {r['code']} | `{r['field']}` | {CHECK[r['check']]} | {param} | {KIND[r['kind']]} "
                f"| {SEVERITY[r['severity']]} | {link} | {cell(r['description'])} |"
            )
        out.append("")
    out += ["## Полные названия источников", ""]
    for key, src in sources.items():
        url = f" — {src['url']}" if src["url"] else ""
        out.append(f"- **{SHORT_SOURCE[key]}**: {src['title']}{url}")
    out += [
        "",
        "## Тестовые документы",
        "",
        "Лежат в [`testdata/documents`](../testdata/documents), ожидаемые поля и замечания —",
        "в [`testdata/expected.json`](../testdata/expected.json).",
        "",
        "| Файл | Тип | Что в нём | Какие правила срабатывают |",
        "|---|---|---|---|",
    ]
    for d in DOCS:
        issues = ", ".join(d["issues"]) or "никакие"
        out.append(f"| `{d['file']}` | `{d['type']}` | {cell(d['about'])} | {issues} |")
    out.append("")
    return "\n".join(out)


# ---------------------------------------------------------------------------
# Запуск
# ---------------------------------------------------------------------------

def generate(testdata_dir: Path, docs_dir: Path) -> None:
    documents = testdata_dir / "documents"
    documents.mkdir(parents=True, exist_ok=True)
    font = register_font()
    for doc in DOCS:
        path = documents / doc["file"]
        if path.suffix == ".docx":
            write_docx(doc, path)
        else:
            write_pdf(doc, path, font)
    (testdata_dir / "expected.json").write_text(expected_json(), encoding="utf-8", newline="\n")
    docs_dir.mkdir(parents=True, exist_ok=True)
    (docs_dir / "RULES.md").write_text(rules_md(), encoding="utf-8", newline="\n")


def same(generated: Path, committed: Path) -> bool:
    return committed.exists() and generated.read_bytes() == committed.read_bytes()


def check() -> None:
    with tempfile.TemporaryDirectory() as tmp:
        generate(Path(tmp) / "testdata", Path(tmp) / "docs")
        pairs = [(Path(tmp) / "docs" / "RULES.md", ROOT / "docs" / "RULES.md"),
                 (Path(tmp) / "testdata" / "expected.json", TESTDATA / "expected.json")]
        pairs += [(Path(tmp) / "testdata" / "documents" / d["file"], TESTDATA / "documents" / d["file"])
                  for d in DOCS]
        stale = [str(committed.relative_to(ROOT)) for generated, committed in pairs
                 if not same(generated, committed)]
        extra = sorted({p.name for p in (TESTDATA / "documents").iterdir()} - {d["file"] for d in DOCS})
    if stale or extra:
        raise GenerationError(
            "Файлы не совпадают с генератором — запустите python testdata/generate.py и закоммитьте:\n  "
            + "\n  ".join(stale + [f"testdata/documents/{name} (лишний файл)" for name in extra])
        )
    print(f"Актуально: {len(DOCS)} документов, expected.json, docs/RULES.md")


def main() -> None:
    try:
        precheck()
        if "--check" in sys.argv[1:]:
            check()
        else:
            generate(TESTDATA, ROOT / "docs")
            print(f"Готово: {len(DOCS)} документов, expected.json, docs/RULES.md")
    except GenerationError as error:
        print(f"Ошибка: {error}", file=sys.stderr)
        sys.exit(1)


if __name__ == "__main__":
    main()
